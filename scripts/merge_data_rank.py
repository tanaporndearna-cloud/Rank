#!/usr/bin/env python3
"""
รวมไฟล์ PPS + PLR + New Product + BarBom + Data Rank เดือนก่อน
เป็น Data Rank รายเดือน (ชีตหลักเป็นสูตร VLOOKUP ล้วน)

ตัวอย่าง:
  python3 merge_data_rank.py \
    --pps PPS08.xlsx --plr PLR08.csv --newproduct NEWPRODUCT08.xlsx \
    --barbom Barbom08.xlsx --prev Data_Rank7.xlsx \
    --month 08 --out DATA_RANK08.xlsx

ต้องมี: pandas, openpyxl  (ถ้ามี LibreOffice จะคำนวณสูตรให้อัตโนมัติ ปิดด้วย --no-recalc)
"""
import argparse
import collections
import os
import re
import shutil
import subprocess
import sys
import tempfile

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill

LIMIT = 30000          # จำนวนแถวสูงสุดที่สูตรครอบคลุม
FILL_LIMIT = 500       # จำนวนแถวสูงสุดของชีตเติมยี่ห้อ/รุ่น
CAT_LIMIT = 3000       # จำนวนแถวสูงสุดของชีต CategoryName
SKIP_BRAND = {"nan", "สินค้า Universal", "สิ่งของในบริษัท"}

OUT_COLS = ["IcCategoryCode", "Product Code", "Description", "ยี่ห้อ", "รุ่น",
            "IcCategoryName", "List Price", "Member", "Wholesale General",
            "Balance (All WH)", "Retail Price", "Bar Bom", "New Product",
            "PropComposer"]


def clean(v):
    return None if (isinstance(v, float) and pd.isna(v)) else v


def norm_code(s):
    return s.astype(str).str.strip()


# ---------------------------------------------------------------- อ่านไฟล์
def read_pps(path):
    df = pd.read_excel(path, dtype={"Product Code": str})
    need = ["Product Code", "Description", "Balance (All WH)", "List Price",
            "Member", "Retail Price", "Wholesale General"]
    miss = [c for c in need if c not in df.columns]
    if miss:
        sys.exit(f"PPS ไม่มีคอลัมน์: {miss}")
    df = df[need].copy()
    df["Product Code"] = norm_code(df["Product Code"])
    df = df[df["Product Code"].ne("nan")]
    dup = df["Product Code"].duplicated().sum()
    if dup:
        print(f"! PPS มีรหัสซ้ำ {dup} แถว (ใช้ตัวแรก)")
        df = df.drop_duplicates("Product Code")
    return df


def read_plr(path):
    use = ["BarCode", "PropComposer", "IcCategoryCode", "IcProductCode"]
    kw = dict(dtype=str, low_memory=False, usecols=lambda c: c in use)
    if path.lower().endswith(".csv"):
        # ERP export ภาษาไทย มักเป็น cp874 และมีไบต์เสียปนอยู่
        for enc in ("utf-8-sig", "cp874"):
            try:
                df = pd.read_csv(path, encoding=enc, **kw)
                break
            except UnicodeDecodeError:
                continue
        else:
            df = pd.read_csv(path, encoding="cp874", encoding_errors="replace", **kw)
    else:
        df = pd.read_excel(path, dtype=str, usecols=lambda c: c in use)
    for c in use:
        if c not in df.columns:
            sys.exit(f"PLR ไม่มีคอลัมน์: {c}")
    df = df[use].copy()
    df["BarCode"] = norm_code(df["BarCode"])
    df["IcProductCode"] = norm_code(df["IcProductCode"])
    return df


def read_codes(path, col_candidates, valid=None):
    """อ่านรายการรหัสสินค้า จาก xlsx หรือ csv (csv ภาษาไทยมักเป็น cp874)
    ไฟล์ New Product บางเดือนเป็น csv รูปแบบเดียวกับ PLR (ใช้คอลัมน์ BarCode / IcProductCode)"""
    if path.lower().endswith(".csv"):
        for enc in ("utf-8-sig", "cp874"):
            try:
                df = pd.read_csv(path, dtype=str, encoding=enc, low_memory=False)
                break
            except UnicodeDecodeError:
                continue
        else:
            df = pd.read_csv(path, dtype=str, encoding="cp874", encoding_errors="replace", low_memory=False)
    else:
        df = pd.read_excel(path, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]
    if valid is not None and "BarCode" in df.columns and "IcProductCode" in df.columns:
        # รูปแบบ PLR: BarCode บางรหัสถูกตัด/ไม่ตรง PPS ให้เลือกตัวที่ตรง PPS (ไม่งั้นใช้ BarCode)
        bc = df["BarCode"].fillna("").astype(str).str.strip()
        ic = df["IcProductCode"].fillna("").astype(str).str.strip()
        codes = pd.Series([b if (b in valid or i not in valid) else i for b, i in zip(bc, ic)])
        codes = codes[(codes != "") & (codes != "รวมทั้งหมด")]
        return list(dict.fromkeys(codes))
    col = next((c for c in col_candidates if c in df.columns), df.columns[0])
    codes = df[col].dropna().astype(str).str.strip()
    codes = codes[(codes != "") & (codes != "รวมทั้งหมด")]
    return list(dict.fromkeys(codes))


def read_prev(path):
    df = pd.read_excel(path, dtype={"Product Code": str})
    df["Product Code"] = norm_code(df["Product Code"])
    # ไฟล์เก่าใช้ชื่อคอลัมน์ Category แทน IcCategoryCode
    if "IcCategoryCode" not in df.columns and "Category" in df.columns:
        df = df.rename(columns={"Category": "IcCategoryCode"})
    for c in ("ยี่ห้อ", "รุ่น"):
        if c not in df.columns:
            sys.exit(f"ไฟล์ Data Rank เดือนก่อนไม่มีคอลัมน์: {c}")
    return df


# ---------------------------------------------------------------- เติมยี่ห้อ/รุ่น
def build_inferer(prev):
    m2b = collections.defaultdict(collections.Counter)
    for b, m in zip(prev["ยี่ห้อ"].astype(str).str.strip().str.upper(),
                    prev["รุ่น"].astype(str).str.strip()):
        if m and m not in SKIP_BRAND:
            m2b[m.upper()][b] += 1
    models = sorted(m2b, key=len, reverse=True)

    def infer(desc):
        d = str(desc).upper().strip()
        if d.startswith("U "):
            return "สินค้า Universal", "สินค้า Universal"
        # 1) ชื่อรุ่นอยู่ต้นรายละเอียด (แม่นที่สุด)
        for m in models:
            if len(m) >= 3 and re.match(r"^" + re.escape(m) + r"(?![A-Z0-9])", d):
                return m2b[m].most_common(1)[0][0], m
        # 2) ชื่อรุ่นอยู่กลางรายละเอียด: เอาตัวที่โผล่ก่อน ถ้าเท่ากันเอาชื่อยาวกว่า
        best = None
        for m in models:
            if len(m) < 3:
                continue
            hit = re.search(r"(?<![A-Z0-9])" + re.escape(m) + r"(?![A-Z0-9])", d)
            if hit and (best is None or hit.start() < best[0]):
                best = (hit.start(), m)
        if best:
            return m2b[best[1]].most_common(1)[0][0], best[1]
        return None, None
    return infer


# ---------------------------------------------------------------- คำนวณค่าสุดท้าย (เท่ากับสูตร VLOOKUP)
def first_map(keys, vals):
    m = {}
    for k, v in zip(keys, vals):
        if k not in m:
            m[k] = v
    return m


def merge_values(pps, plr, new_codes, bb_codes, prev):
    """คืน (rows, fills)  rows = แถวตาม OUT_COLS ค่าเท่ากับสูตร VLOOKUP ในชีตหลัก
    fills = [(code, ยี่ห้อ, รุ่น, description)] ที่เติมเอง"""
    cat_by_bar = first_map(plr["BarCode"], plr["IcCategoryCode"])
    cat_by_ic = first_map(plr["IcProductCode"], plr["IcCategoryCode"])
    pc_by_bar = first_map(plr["BarCode"], plr["PropComposer"])
    pc_by_ic = first_map(plr["IcProductCode"], plr["PropComposer"])
    prev_brand = first_map(prev["Product Code"], prev["ยี่ห้อ"])
    prev_model = first_map(prev["Product Code"], prev["รุ่น"])
    if "IcCategoryName" in prev.columns and "IcCategoryCode" in prev.columns:
        cn = prev[["IcCategoryCode", "IcCategoryName"]].dropna().drop_duplicates("IcCategoryCode")
        cat_name = dict(zip(cn["IcCategoryCode"], cn["IcCategoryName"]))
    else:
        cat_name = {}
    infer = build_inferer(prev)
    fill = {}
    for code, desc in zip(pps["Product Code"], pps["Description"]):
        if code not in prev_brand:
            b, m = infer(desc)
            if b:
                fill[code] = (b, m, desc)
    newset, bbset = set(new_codes), set(bb_codes)

    def nz(v):                      # VLOOKUP ช่องว่างคืน 0
        return 0 if (v is None or (isinstance(v, float) and pd.isna(v))) else v

    def txt(v):
        return "" if (v is None or (isinstance(v, float) and pd.isna(v))) else str(v)

    rows = []
    for t in pps.itertuples(index=False):
        code = t[0]
        if code in cat_by_bar:
            cat = txt(cat_by_bar[code])
        elif code in cat_by_ic:
            cat = txt(cat_by_ic[code])
        else:
            cat = ""
        if code in prev_brand:
            brand, model = nz(prev_brand[code]), nz(prev_model[code])
        elif code in fill:
            brand, model = fill[code][0], fill[code][1]
        else:
            brand = model = None                      # #N/A ต้องเติมมือ
        pc = pc_by_bar.get(code) if code in pc_by_bar else pc_by_ic.get(code, "")
        rows.append([cat, code, txt(t[1]), brand, model,
                     cat_name.get(cat, "") if cat != "" else "",
                     nz(t[3]), nz(t[4]), nz(t[6]), nz(t[2]), nz(t[5]),
                     "Bar Bom" if code in bbset else "", "New Product" if code in newset else "", txt(pc)])
    fills = [(c, v[0], v[1], v[2]) for c, v in fill.items()]
    return rows, fills


def summarize_rows(rows, pps, new_codes, bb_codes):
    ix = {c: i for i, c in enumerate(OUT_COLS)}
    blank = lambda v: v is None or v == ""
    need = [r[ix["Product Code"]] for r in rows if blank(r[ix["ยี่ห้อ"]]) or blank(r[ix["รุ่น"]])]
    print("\n=== สรุป ===")
    print(f"แถวทั้งหมด        : {len(rows):,} (PPS {len(pps):,})")
    print(f"New Product       : {sum(not blank(r[ix['New Product']]) for r in rows)} / ในไฟล์ {len(new_codes)}")
    print(f"Bar Bom           : {sum(not blank(r[ix['Bar Bom']]) for r in rows)} / ในไฟล์ {len(bb_codes)}")
    print(f"IcCategoryCode ว่าง: {sum(blank(r[ix['IcCategoryCode']]) for r in rows)}")
    print(f"ยี่ห้อ/รุ่น ต้องเติมมือ: {len(need)}")
    for code in need[:20]:
        print("   -", code)
    miss_new = set(new_codes) - set(pps["Product Code"])
    miss_bb = set(bb_codes) - set(pps["Product Code"])
    if miss_new:
        print(f"! New Product ไม่อยู่ใน PPS: {sorted(miss_new)[:10]}")
    if miss_bb:
        print(f"! BarBom ไม่อยู่ใน PPS: {sorted(miss_bb)[:10]}")


# ---------------------------------------------------------------- เขียนสมุดงาน
def write_sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append([clean(v) for v in r])
    ws.freeze_panes = "A2"
    for c in ws[1]:
        c.font = Font(bold=True)
    return ws


def build_workbook(pps, plr, new_codes, bb_codes, prev, month, out_path):
    wb = Workbook()
    main = wb.active
    main.title = f"Data Rank{month}"

    write_sheet(wb, "PPS", list(pps.columns), pps.itertuples(index=False))
    write_sheet(wb, "PLR", ["BarCode", "PropComposer", "IcCategoryCode", "IcProductCode"],
                plr[["BarCode", "PropComposer", "IcCategoryCode", "IcProductCode"]].itertuples(index=False))
    write_sheet(wb, "New Product", ["BarCode", "New Product"], [(c, "New Product") for c in new_codes])
    write_sheet(wb, "Bar Bom", ["Product Code", "Bar Bom"], [(c, "Bar Bom") for c in bb_codes])
    write_sheet(wb, "Rank เดือนก่อน", ["Product Code", "ยี่ห้อ", "รุ่น"],
                prev[["Product Code", "ยี่ห้อ", "รุ่น"]].itertuples(index=False))

    if "IcCategoryName" in prev.columns and "IcCategoryCode" in prev.columns:
        cat = prev[["IcCategoryCode", "IcCategoryName"]].dropna().drop_duplicates("IcCategoryCode")
    else:
        cat = pd.DataFrame(columns=["IcCategoryCode", "IcCategoryName"])
    write_sheet(wb, "CategoryName", ["IcCategoryCode", "IcCategoryName"], cat.itertuples(index=False))

    # หายี่ห้อ/รุ่นที่เดือนก่อนไม่มี แล้วเติมจากรายละเอียด
    prev_codes = set(prev["Product Code"])
    infer = build_inferer(prev)
    fills = []
    for code, desc in zip(pps["Product Code"], pps["Description"]):
        if code not in prev_codes:
            b, m = infer(desc)
            if b:
                fills.append((code, b, m, desc))
    ws_fill = write_sheet(wb, "เติมยี่ห้อรุ่น (auto)", ["Product Code", "ยี่ห้อ", "รุ่น", "Description"], fills)
    n_fill = len(fills)

    # ชีตหลัก
    main.append(OUT_COLS)
    P = f"PPS!$A$2:$G${LIMIT}"
    PREV = f"'Rank เดือนก่อน'!$A$2:$C${LIMIT}"
    FILL = f"'เติมยี่ห้อรุ่น (auto)'!$A$2:$C${FILL_LIMIT}"
    for i, code in enumerate(pps["Product Code"], start=2):
        main.append([
            f'=IFERROR(VLOOKUP(B{i},PLR!$A$2:$C${LIMIT},3,0),IFERROR(INDEX(PLR!$C$2:$C${LIMIT},MATCH(B{i},PLR!$D$2:$D${LIMIT},0)),""))',
            code,
            f'=IFERROR(VLOOKUP(B{i},{P},2,0)&"","")',
            f'=IFERROR(VLOOKUP(B{i},{PREV},2,0),VLOOKUP(B{i},{FILL},2,0))',
            f'=IFERROR(VLOOKUP(B{i},{PREV},3,0),VLOOKUP(B{i},{FILL},3,0))',
            f'=IFERROR(VLOOKUP(A{i},CategoryName!$A$2:$B${CAT_LIMIT},2,0),"")',
            f'=IFERROR(VLOOKUP(B{i},{P},4,0),"")',
            f'=IFERROR(VLOOKUP(B{i},{P},5,0),"")',
            f'=IFERROR(VLOOKUP(B{i},{P},7,0),"")',
            f'=IFERROR(VLOOKUP(B{i},{P},3,0),"")',
            f'=IFERROR(VLOOKUP(B{i},{P},6,0),"")',
            f"=IFERROR(VLOOKUP(B{i},'Bar Bom'!$A$2:$B${LIMIT},2,0),\"\")",
            f"=IFERROR(VLOOKUP(B{i},'New Product'!$A$2:$B${LIMIT},2,0),\"\")",
            f'=IFERROR(VLOOKUP(B{i},PLR!$A$2:$B${LIMIT},2,0)&"",IFERROR(INDEX(PLR!$B$2:$B${LIMIT},MATCH(B{i},PLR!$D$2:$D${LIMIT},0))&"",""))',
        ])
    main.freeze_panes = "A2"
    for c in main[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEFGHIJKLMN", [14, 20, 60, 18, 18, 22, 10, 10, 12, 12, 10, 9, 12, 12]):
        main.column_dimensions[col].width = w
    yellow = PatternFill("solid", bgColor="FFFF00", fgColor="FFFF00")
    n = len(pps)
    main.conditional_formatting.add(
        f"D2:E{n + 1}",
        FormulaRule(formula=[f"ISNA(MATCH($B2,'Rank เดือนก่อน'!$A$2:$A${LIMIT},0))"], fill=yellow))
    # ชีตค่า (ไม่ใช่สูตร) ใช้เป็นอินพุตของขั้นจัด Rank ได้ทันที ไม่ต้องเปิดคำนวณใน Excel
    rows, _fills = merge_values(pps, plr, new_codes, bb_codes, prev)
    ws_val = wb.create_sheet(f"Data Rank{month} (ค่า)")
    ws_val.append(OUT_COLS)
    for r in rows:
        ws_val.append([clean(v) for v in r])
    ws_val.freeze_panes = "A2"
    for c in ws_val[1]:
        c.font = Font(bold=True)
    wb.save(out_path)
    return n_fill, rows


# ---------------------------------------------------------------- คำนวณสูตร + ตรวจ
def recalc(path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("! ไม่พบ LibreOffice ข้ามการคำนวณสูตร (Excel จะคำนวณเองตอนเปิดไฟล์)")
        return False
    tmp = tempfile.mkdtemp()
    subprocess.run([soffice, "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                    "--outdir", tmp, path], check=True, capture_output=True, timeout=900)
    shutil.move(os.path.join(tmp, os.path.basename(path)), path)
    return True


def summarize(path, month, pps, new_codes, bb_codes):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True)[f"Data Rank{month}"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    ix = {c: i for i, c in enumerate(OUT_COLS)}
    blank = lambda v: v is None or v == ""
    need = [r[ix["Product Code"]] for r in rows
            if blank(r[ix["ยี่ห้อ"]]) or blank(r[ix["รุ่น"]])
            or str(r[ix["ยี่ห้อ"]]).startswith("#") or str(r[ix["รุ่น"]]).startswith("#")]
    print("\n=== สรุป ===")
    print(f"แถวทั้งหมด        : {len(rows):,} (PPS {len(pps):,})")
    print(f"New Product       : {sum(not blank(r[ix['New Product']]) for r in rows)} / ในไฟล์ {len(new_codes)}")
    print(f"Bar Bom           : {sum(not blank(r[ix['Bar Bom']]) for r in rows)} / ในไฟล์ {len(bb_codes)}")
    print(f"IcCategoryCode ว่าง: {sum(blank(r[ix['IcCategoryCode']]) for r in rows)}")
    print(f"ยี่ห้อ/รุ่น ต้องเติมมือ: {len(need)}")
    for code in need[:20]:
        print("   -", code)
    miss_new = set(new_codes) - set(pps["Product Code"])
    miss_bb = set(bb_codes) - set(pps["Product Code"])
    if miss_new:
        print(f"! New Product ไม่อยู่ใน PPS: {sorted(miss_new)[:10]}")
    if miss_bb:
        print(f"! BarBom ไม่อยู่ใน PPS: {sorted(miss_bb)[:10]}")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--pps", required=True)
    ap.add_argument("--plr", required=True)
    ap.add_argument("--newproduct", required=True)
    ap.add_argument("--barbom", required=True)
    ap.add_argument("--prev", required=True, help="Data Rank เดือนก่อน (ใช้ ยี่ห้อ/รุ่น/ชื่อประเภท)")
    ap.add_argument("--month", required=True, help="เลขเดือน เช่น 08")
    ap.add_argument("--out", required=True)
    ap.add_argument("--no-recalc", action="store_true")
    a = ap.parse_args(argv)

    pps = read_pps(a.pps)
    plr = read_plr(a.plr)
    new_codes = read_codes(a.newproduct, ["รหัสสินค้า", "Product Code", "BarCode", "IcProductCode"], valid=set(pps["Product Code"]))
    bb_codes = read_codes(a.barbom, ["Product Code", "รหัสสินค้า"])
    prev = read_prev(a.prev)
    print(f"PPS {len(pps):,} | PLR {len(plr):,} | New {len(new_codes)} | BarBom {len(bb_codes)} | เดือนก่อน {len(prev):,}")
    if len(pps) > LIMIT - 1:
        sys.exit(f"แถวเกิน {LIMIT} ต้องปรับค่า LIMIT ในสคริปต์")

    n_fill, rows = build_workbook(pps, plr, new_codes, bb_codes, prev, a.month, a.out)
    print(f"สร้างไฟล์แล้ว: {a.out} | เติมยี่ห้อ/รุ่นอัตโนมัติ {n_fill} รายการ")
    summarize_rows(rows, pps, new_codes, bb_codes)
    if not a.no_recalc:
        recalc(a.out)


if __name__ == "__main__":
    main()
